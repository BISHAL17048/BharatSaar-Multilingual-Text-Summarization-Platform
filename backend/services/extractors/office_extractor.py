from services.extractors.base import BaseExtractor
import docx
from pptx import Presentation

class OfficeExtractor(BaseExtractor):
    def __init__(self, ext: str):
        self.ext = ext

    async def extract(self, file_path: str) -> dict:
        if self.ext == "DOCX":
            return self._extract_docx(file_path)
        elif self.ext == "PPTX":
            return self._extract_pptx(file_path)
        raise ValueError("Unsupported office format")

    def _extract_docx(self, file_path: str) -> dict:
        doc = docx.Document(file_path)
        full_text = []
        tables_content = []
        
        for para in doc.paragraphs:
            full_text.append(para.text)
            
        for table in doc.tables:
            # Simple markdown table conversion
            for row in table.rows:
                row_text = "| " + " | ".join(cell.text.replace("\n", " ") for cell in row.cells) + " |"
                tables_content.append(row_text)
                
        return {
            "raw_text": "\n".join(full_text),
            "tables": tables_content,
            "metadata": {"author": doc.core_properties.author}
        }

    def _extract_pptx(self, file_path: str) -> dict:
        prs = Presentation(file_path)
        full_text = []
        
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    full_text.append(shape.text)
                    
        return {
            "raw_text": "\n".join(full_text),
            "tables": [],
            "metadata": {"author": prs.core_properties.author}
        }
